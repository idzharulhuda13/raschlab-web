"""F15: the S.E. filter on the Bandingkan view (PLAN_F15.md).

Five checks over the frozen contract:
1. An item whose measure moved 0,10 stays out while the one that moved 0,40
   stays in, because an item's noise floor is its own S.E. (0,15 here).
2. The live count line names the kept and the total counts when the filter is
   on, and only the total when it is off.
3. An item whose S.E. cell is empty is never reported as moved.
4. The export sheet filters the same way without moving its header row or its
   five columns.
5. The compare fragment renders no em dash.

The analysis rows are seeded the same way tests/test_explorer_routes.py seeds
them (no engine run); only the item tables change, to carry the S.E. column.
"""

from __future__ import annotations

import io

import openpyxl
from fastapi.testclient import TestClient

from app.export import COMPARE_HEADER
from tests.test_analysis_routes import _create_authenticated_user
from tests.test_explorer_routes import (
    _ITEM_HEADER_ROW0,
    _ITEM_HEADER_ROW1,
    _csv,
    _seed_done,
)

_SE = "0.15"


def _item_row(entry: int, measure: str, se: str, label: str) -> list[str]:
    """A 14-column item row shaped like the engine's, with a chosen S.E."""
    return [
        str(entry), "50", "30", measure, se,
        "1.10", "0.50", "1.05", "0.30", "0.55", "0.52", "73.5", "72.0", label,
    ]


def _item_csv(rows: list[list[str]]) -> str:
    return _csv([_ITEM_HEADER_ROW0, _ITEM_HEADER_ROW1] + rows)


def _seed_pair(
    client: TestClient,
    email: str,
    rows_from: list[list[str]],
    rows_to: list[list[str]],
) -> tuple[int, int]:
    """Seed two finished analyses, one item table each, and return their ids."""
    user_id = _create_authenticated_user(client, email=email)
    id_from = _seed_done(
        user_id, "se_pertama.csv", {"item_table_15.1.csv": _item_csv(rows_from)}
    )
    id_to = _seed_done(
        user_id, "se_kedua.csv", {"item_table_15.1.csv": _item_csv(rows_to)}
    )
    return id_from, id_to


def _frag_url(analysis_id: int, id_from: int, id_to: int, *, over_se: bool = False) -> str:
    url = (
        f"/analyses/{analysis_id}/explore?view=bandingkan&fragment=bandingkan"
        f"&from={id_from}&to={id_to}"
    )
    if over_se:
        url += "&over_se=1"
    return url


# Two items move 0,10 and one moves 0,40; every S.E. is 0,15.
_ROWS_FROM = [
    _item_row(1, "0.00", _SE, "I1"),
    _item_row(2, "0.00", _SE, "I2"),
    _item_row(3, "0.00", _SE, "I3"),
]
_ROWS_TO = [
    _item_row(1, "0.10", _SE, "I1"),
    _item_row(2, "0.10", _SE, "I2"),
    _item_row(3, "0.40", _SE, "I3"),
]


def test_only_items_that_moved_more_than_their_se(client: TestClient):
    id_from, id_to = _seed_pair(client, "over_se_rows@example.test", _ROWS_FROM, _ROWS_TO)

    unfiltered = client.get(_frag_url(id_from, id_from, id_to))
    assert unfiltered.status_code == 200
    assert unfiltered.text.count('class="delta-row"') == 3
    assert "+0.10" in unfiltered.text
    assert "+0.40" in unfiltered.text

    filtered = client.get(_frag_url(id_from, id_from, id_to, over_se=True))
    assert filtered.status_code == 200
    assert filtered.text.count('class="delta-row"') == 1
    assert "+0.40" in filtered.text
    assert "+0.10" not in filtered.text


def test_live_count_line_names_kept_and_total(client: TestClient):
    id_from, id_to = _seed_pair(client, "over_se_count@example.test", _ROWS_FROM, _ROWS_TO)

    filtered = client.get(_frag_url(id_from, id_from, id_to, over_se=True))
    assert filtered.status_code == 200
    assert "Menampilkan 1 dari 3 butir, yang bergeser lebih dari S.E.-nya." in filtered.text

    unfiltered = client.get(_frag_url(id_from, id_from, id_to))
    assert unfiltered.status_code == 200
    assert "Menampilkan 3 butir." in unfiltered.text


def test_item_without_se_is_never_reported_as_moved(client: TestClient):
    rows_from = [
        _item_row(1, "0.00", "", "I1"),
        _item_row(2, "0.00", _SE, "I2"),
    ]
    rows_to = [
        _item_row(1, "0.40", "", "I1"),
        _item_row(2, "0.10", _SE, "I2"),
    ]
    id_from, id_to = _seed_pair(client, "over_se_missing@example.test", rows_from, rows_to)

    unfiltered = client.get(_frag_url(id_from, id_from, id_to))
    assert unfiltered.status_code == 200
    assert unfiltered.text.count('class="delta-row"') == 2
    assert "+0.40" in unfiltered.text

    filtered = client.get(_frag_url(id_from, id_from, id_to, over_se=True))
    assert filtered.status_code == 200
    assert filtered.text.count('class="delta-row"') == 0
    assert "+0.40" not in filtered.text
    assert "Menampilkan 0 dari 2 butir, yang bergeser lebih dari S.E.-nya." in filtered.text
    assert (
        "Tidak ada butir yang bergeser lebih dari S.E.-nya. "
        "Matikan filter untuk melihat semua butir."
    ) in filtered.text


def test_export_over_se_keeps_header_and_reduces_rows(client: TestClient):
    id_from, id_to = _seed_pair(client, "over_se_export@example.test", _ROWS_FROM, _ROWS_TO)

    base = f"/analyses/{id_from}/export?table=bandingkan&from={id_from}&to={id_to}"

    full = client.get(base)
    assert full.status_code == 200
    ws_full = openpyxl.load_workbook(io.BytesIO(full.content), data_only=False)["bandingkan"]

    filtered = client.get(base + "&over_se=1")
    assert filtered.status_code == 200
    ws_filtered = openpyxl.load_workbook(io.BytesIO(filtered.content), data_only=False)["bandingkan"]

    header_full = [cell.value for cell in ws_full[1]]
    header_filtered = [cell.value for cell in ws_filtered[1]]
    assert header_full == COMPARE_HEADER
    assert header_filtered == header_full

    assert ws_full.max_row == 4
    assert ws_filtered.max_row == 2

    kept_row = [cell.value for cell in ws_filtered[2]]
    assert int(kept_row[0]) == 3
    assert float(kept_row[4]) == 0.40


def test_compare_fragment_has_no_em_dash(client: TestClient):
    id_from, id_to = _seed_pair(client, "over_se_dash@example.test", _ROWS_FROM, _ROWS_TO)

    html = client.get(_frag_url(id_from, id_from, id_to, over_se=True))
    assert html.status_code == 200
    assert "\u2014" not in html.text
