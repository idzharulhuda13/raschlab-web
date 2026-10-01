"""Tests for the combined workbook export (/analyses/{id}/export?table=semua)."""

import gzip
import hashlib
import io
import json
from typing import Optional
from fastapi.testclient import TestClient
import openpyxl
import pytest

from app.auth import COOKIE_NAME
from app.db import SessionLocal
from app.export import (
    EXPORT_MAX_CELLS,
    XLSX_MEDIA_TYPE,
)
from app.models import Analysis, AnalysisFile, Dataset, SessionRow, User
from app.security import hash_password, hash_token, new_token, now_epoch


def _csv(rows: list[list[str]]) -> str:
    """Produce CSV text: commas, CRLF, trailing CRLF."""
    return "\r\n".join(",".join(cell for cell in row) for row in rows) + "\r\n"


_ITEM_ROWS = [
    ["ENTRY", "TOTAL", "TOTAL", "JMLE", "MODEL", "INFIT", "INFIT", "OUTFIT", "OUTFIT", "PTMEASUR-AL", "EXP.", "EXACT", "EXACT", "ITEM"],
    ["", "SCORE", "COUNT", "MEASURE", "S.E.", "MNSQ", "ZSTD", "MNSQ", "ZSTD", "CORR.", "", "OBS%", "EXPECTED%", ""],
    ["1", "10", "20", "0.10", "0.20", "1.00", "0.0", "1.00", "0.0", "0.50", "0.45", "75.0", "70.0", "Item1"],
    ["2", "15", "20", "-0.25", "0.22", "0.95", "-0.2", "0.98", "-0.1", "0.45", "0.40", "80.0", "72.0", "Item2"],
]

_PERSON_ROWS = [
    ["ENTRY", "TOTAL", "TOTAL", "JMLE", "MODEL", "INFIT", "INFIT", "OUTFIT", "OUTFIT", "PTMEASUR-AL", "EXP.", "EXACT", "EXACT", "PERSON"],
    ["", "SCORE", "COUNT", "MEASURE", "S.E.", "MNSQ", "ZSTD", "MNSQ", "ZSTD", "CORR.", "", "OBS%", "EXPECTED%", ""],
    ["1", "10", "12", "0.50", "0.20", "1.00", "0.0", "1.00", "0.0", "0.40", "0.35", "70.0", "65.0", "Person_1"],
    ["2", "11", "12", "0.60", "0.21", "0.95", "-0.1", "0.98", "-0.1", "0.42", "0.36", "75.0", "68.0", "Person_2"],
]

_OPTION_ROWS = [
    ["ENTRY", "DATA", "SCORE", "DATA", "", "ABILITY", "", "S.E.", "INFT", "OUTF", "PTMA", ""],
    ["NUMBER", "CODE", "VALUE", "COUNT", "%", "ABILITY MEAN", "ABILITY PSD", "SE MEAN", "INFT MNSQ", "OUTF MNSQ", "PTMA CORR", "ITEM"],
    ["1", "0", "0", "5", "25.0", "-0.40", "0.30", "0.20", "1.05", "1.02", "-0.30", "Item1"],
    ["1", "1", "1", "15", "75.0", "0.50", "0.35", "0.21", "0.98", "0.96", "0.50", "Item1"],
]

_SUMMARY_ROWS = [
    ["SECTION", "STATISTIC", "VALUE"],
    ["", "", ""],
    ["ITEM", "COUNT", "2"],
    ["ITEM", "MEASURE MEAN", "-0.08"],
    ["PERSON", "COUNT", "2"],
    ["PERSON", "MEASURE MEAN", "0.55"],
]

_WRIGHT_MEASURE_ROWS = [
    ["MEASURE", "NR_PERSON", "PERSON_HIST", "NR_ITEM", "ITEMS", "ITEM_HIST", "PERSON_ENTRIES", "ITEM_ENTRIES"],
    ["", "", "", "", "", "", "PERSON_ENTRIES", "ITEM_ENTRIES"],
    ["1.00", "1", "#", "0", "", "", "2", ""],
    ["0.50", "1", "#", "0", "", "", "1", ""],
    ["0.00", "0", "", "1", "Item1", "#", "", "1"],
    ["-0.50", "0", "", "1", "Item2", "#", "", "2"],
]

_WRIGHT_FREQ_ROWS = [
    ["MEASURE", "NR_PERSON", "NR_PERSON_PRESENT", "PERSON_HIST", "PERSON_FREQ_HIST", "NR_ITEM", "ITEMS", "ITEM_HIST", "PERSON_ENTRIES", "ITEM_ENTRIES"],
    ["", "", "", "", "", "", "", "", "PERSON_ENTRIES", "ITEM_ENTRIES"],
    ["1.00", "1", "1", "#", "#", "0", "", "", "2", ""],
    ["0.50", "1", "1", "#", "#", "0", "", "", "1", ""],
    ["0.00", "0", "0", "", "", "1", "Item1", "#", "", "1"],
    ["-0.50", "0", "0", "", "", "1", "Item2", "#", "", "2"],
]


def _seed_done(
    user_id: int,
    filename: str,
    files: dict[str, str],
    status: str = "done",
    created_at: Optional[int] = None,
    n_persons: int = 2,
    n_items: int = 2,
) -> int:
    now = created_at if created_at is not None else now_epoch()
    with SessionLocal() as db:
        raw_text = b"dummy,csv,content"
        raw_gz = gzip.compress(raw_text, mtime=0)
        item_labels = [f"Item{i+1}" for i in range(n_items)]
        dataset = Dataset(
            user_id=user_id,
            filename=filename,
            kind="delimited",
            format="csv",
            status="ready",
            n_persons=n_persons,
            n_items=n_items,
            item_labels_json=json.dumps(item_labels),
            mapping_json=json.dumps({"1": "correct", "0": "incorrect"}),
            summary_json=json.dumps({}),
            raw_gzip=raw_gz,
            raw_bytes=len(raw_text),
            created_at=now,
            committed_at=now,
        )
        db.add(dataset)
        db.commit()
        dataset_id = dataset.id

        analysis = Analysis(
            user_id=user_id,
            dataset_id=dataset_id,
            status=status,
            params_json=json.dumps({"mode": "compat", "digits": 2}),
            engine_ref="test-ref",
            created_at=now,
            expires_at=now + 180 * 86400,
            finished_at=now if status == "done" else None,
            elapsed_ms=1234 if status == "done" else None,
        )
        db.add(analysis)
        db.commit()
        analysis_id = analysis.id

        for fname, text in files.items():
            raw_bytes = text.encode("utf-8")
            af = AnalysisFile(
                analysis_id=analysis_id,
                filename=fname,
                content_gzip=gzip.compress(raw_bytes, mtime=0),
                sha256=hashlib.sha256(raw_bytes).hexdigest(),
                bytes=len(raw_bytes),
            )
            db.add(af)
        db.commit()

    return analysis_id


def _create_user_with_token(client: TestClient, email: str = "wb_user@example.test") -> tuple[int, str]:
    with SessionLocal() as db:
        now = now_epoch()
        user = User(
            email=email,
            email_normalized=email.lower().strip(),
            password_hash=hash_password("katasandi-wb-123"),
            created_at=now,
            verified_at=now,
        )
        db.add(user)
        db.commit()
        user_id = user.id

        token = new_token()
        session_row = SessionRow(
            user_id=user_id,
            token_hash=hash_token(token),
            created_at=now,
            expires_at=now + 86400 * 30,
        )
        db.add(session_row)
        db.commit()

    client.cookies.set(COOKIE_NAME, token)
    return user_id, token


def test_combined_export_completed_analysis(client):
    user_id, token = _create_user_with_token(client, email="wb_export@example.test")
    client.cookies.set(COOKIE_NAME, token)

    files = {
        "item_table_15.1.csv": _csv(_ITEM_ROWS),
        "person_table.csv": _csv(_PERSON_ROWS),
        "option_table_15.3.csv": _csv(_OPTION_ROWS),
        "summary_table.csv": _csv(_SUMMARY_ROWS),
        "wright_map_measure.csv": _csv(_WRIGHT_MEASURE_ROWS),
        "wright_map_frequency.csv": _csv(_WRIGHT_FREQ_ROWS),
    }
    analysis_id = _seed_done(user_id=user_id, filename="data_wb.csv", files=files)

    resp = client.get(f"/analyses/{analysis_id}/export?table=semua")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_MEDIA_TYPE

    expected_filename = f"raschlab-{analysis_id}-semua.xlsx"
    cd = resp.headers["content-disposition"]
    assert f'filename="{expected_filename}"' in cd
    assert f"filename*=UTF-8''{expected_filename}" in cd
    assert resp.headers["cache-control"] == "no-store"

    # Assert that bytes open with openpyxl.load_workbook
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)

    # Assert exact six sheet names in order
    assert wb.sheetnames == [
        "15.1",
        "person",
        "15.3",
        "summary",
        "wright_measure",
        "wright_frequency",
    ]

    # Assert at least one numeric cell in the 15.1 sheet is a real number (int or float)
    ws_item = wb["15.1"]
    has_numeric = False
    for r in range(3, ws_item.max_row + 1):
        for c in range(1, ws_item.max_column + 1):
            val = ws_item.cell(row=r, column=c).value
            if isinstance(val, (int, float)):
                has_numeric = True
                break
        if has_numeric:
            break
    assert has_numeric, "Expected at least one numeric cell (int or float) in sheet 15.1"


def test_unknown_table_answers_404(client):
    user_id, token = _create_user_with_token(client, email="wb_unknown@example.test")
    client.cookies.set(COOKIE_NAME, token)

    files = {
        "item_table_15.1.csv": _csv(_ITEM_ROWS),
    }
    analysis_id = _seed_done(user_id=user_id, filename="data_wb_unknown.csv", files=files)

    resp = client.get(f"/analyses/{analysis_id}/export?table=unknown_table")
    assert resp.status_code == 404

    resp_empty = client.get(f"/analyses/{analysis_id}/export?table=")
    assert resp_empty.status_code == 404


def test_combined_workbook_guard_counts_the_subsubtes_sheet(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    user_id, token = _create_user_with_token(client, email="wb_subsubtes_guard@example.test")
    client.cookies.set(COOKIE_NAME, token)

    subsubtes_rows = [
        ["SUBSUBTES", "ITEMS", "ANCHOR_ITEMS", "NEW_ITEMS", "MEAN_MEASURE", "S.SD_MEASURE", "MEAN_INFIT", "MAX_INFIT", "MISFIT_ITEMS"],
        ["Deretan Bilangan", "20", "0", "20", "0.42", "0.15", "1.02", "1.60", "1"],
        ["Logis", "20", "20", "0", "-0.50", "0.10", "0.95", "0.98", "0"],
    ]
    files = {
        "item_table_15.1.csv": _csv(_ITEM_ROWS),
        "person_table.csv": _csv(_PERSON_ROWS),
        "option_table_15.3.csv": _csv(_OPTION_ROWS),
        "summary_table.csv": _csv(_SUMMARY_ROWS),
        "wright_map_measure.csv": _csv(_WRIGHT_MEASURE_ROWS),
        "wright_map_frequency.csv": _csv(_WRIGHT_FREQ_ROWS),
        "subsubtes_summary.csv": _csv(subsubtes_rows),
    }
    analysis_id = _seed_done(user_id=user_id, filename="data_wb_subsubtes_guard.csv", files=files)

    total_rows = (
        len(_ITEM_ROWS)
        + len(_PERSON_ROWS)
        + len(_OPTION_ROWS)
        + len(_SUMMARY_ROWS)
        + len(_WRIGHT_MEASURE_ROWS)
        + len(_WRIGHT_FREQ_ROWS)
        + len(subsubtes_rows)
    )
    total_cells = (
        sum(len(r) for r in _ITEM_ROWS)
        + sum(len(r) for r in _PERSON_ROWS)
        + sum(len(r) for r in _OPTION_ROWS)
        + sum(len(r) for r in _SUMMARY_ROWS)
        + sum(len(r) for r in _WRIGHT_MEASURE_ROWS)
        + sum(len(r) for r in _WRIGHT_FREQ_ROWS)
        + sum(len(r) for r in subsubtes_rows)
    )

    monkeypatch.setattr("app.export.EXPORT_MAX_CELLS", total_cells - 1)
    resp_413 = client.get(f"/analyses/{analysis_id}/export?table=semua")
    assert resp_413.status_code == 413

    monkeypatch.setattr("app.export.EXPORT_MAX_CELLS", total_cells)
    resp_200 = client.get(f"/analyses/{analysis_id}/export?table=semua")
    assert resp_200.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp_200.content))
    assert "subsubtes" in wb.sheetnames


def test_combined_export_caps_summed(client, monkeypatch):
    user_id, token = _create_user_with_token(client, email="wb_caps@example.test")
    client.cookies.set(COOKIE_NAME, token)

    files = {
        "item_table_15.1.csv": _csv(_ITEM_ROWS),
        "person_table.csv": _csv(_PERSON_ROWS),
        "option_table_15.3.csv": _csv(_OPTION_ROWS),
        "summary_table.csv": _csv(_SUMMARY_ROWS),
        "wright_map_measure.csv": _csv(_WRIGHT_MEASURE_ROWS),
        "wright_map_frequency.csv": _csv(_WRIGHT_FREQ_ROWS),
    }
    analysis_id = _seed_done(user_id=user_id, filename="data_wb_caps.csv", files=files)

    total_rows = len(_ITEM_ROWS) + len(_PERSON_ROWS) + len(_OPTION_ROWS) + len(_SUMMARY_ROWS) + len(_WRIGHT_MEASURE_ROWS) + len(_WRIGHT_FREQ_ROWS)
    total_cells = sum(len(r) for r in _ITEM_ROWS) + sum(len(r) for r in _PERSON_ROWS) + sum(len(r) for r in _OPTION_ROWS) + sum(len(r) for r in _SUMMARY_ROWS) + sum(len(r) for r in _WRIGHT_MEASURE_ROWS) + sum(len(r) for r in _WRIGHT_FREQ_ROWS)

    monkeypatch.setattr("app.export.EXPORT_MAX_CELLS", 10)
    resp_413 = client.get(f"/analyses/{analysis_id}/export?table=semua")
    assert resp_413.status_code == 413
    assert str(total_rows) in resp_413.text
    assert str(total_cells) in resp_413.text
