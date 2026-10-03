"""Tests for Ringkasan Tabulasi band and tabulasi_summary.csv / tabulasi_item.csv."""

import io
from io import BytesIO
from fastapi.testclient import TestClient
import openpyxl
from sqlalchemy import select

from app.db import SessionLocal
from app.export import XLSX_MEDIA_TYPE
from app.models import AnalysisFile
from tests.test_analysis_routes import FIXTURES_DIR, _create_authenticated_user
from tests.test_explorer_routes import (
    _ITEM_HEADER_ROW0,
    _ITEM_HEADER_ROW1,
    _PERSON_HEADER_ROW0,
    _PERSON_HEADER_ROW1,
    _WRIGHT_HEADER_ROW0,
    _WRIGHT_HEADER_ROW1,
    _csv,
    _make_item_row,
    _seed_done,
)

_TABULASI_HEADER = ["SUBTES", "SUBSUBTES", "KESUKARAN", "TINGGI", "NOMOR_TINGGI", "RENDAH", "NOMOR_RENDAH", "JUMLAH"]
_TABULASI_ROWS = [
    ["Kuantitatif", "Aritmatika dan Aljabar", "sulit", "1234", "51, 54, 55", "3", "92, 93, 94", "1237"],
    ["Penalaran", "Logis", "mudah", "5", "1, 2", "0", "", "5"],
]
_TABULASI_ITEM_HEADER = ["SUBTES", "SUBSUBTES", "ENTRY", "ITEM", "KESUKARAN", "DAYA_BEDA", "DATA_PCT", "PTMA_CORR"]
_TABULASI_ITEM_ROWS = [
    ["Kuantitatif", "Aritmatika dan Aljabar", "1", "01tbskda26a01", "sulit", "rendah", "21", "0.10"],
    ["Kuantitatif", "Aritmatika dan Aljabar", "2", "01tbskda26a02", "mudah", "tinggi", "70", "0.45"],
]

_OLD_RUN_NOTE = (
    "Analisis ini dijalankan sebelum mesin menulis ringkasan tabulasi, "
    "jadi rinciannya belum tersimpan. Jalankan ulang analisis untuk melihatnya."
)
_HEADER_ONLY_NOTE = (
    "Data ini tidak memuat kode sub-subtes pada label butir, "
    "jadi ringkasan tabulasi kosong. Analisis tetap lengkap; "
    "tabel ini terisi otomatis bila label butir memuat kode sub-subtes."
)


def _csv_table(rows: list[list[str]]) -> str:
    return _csv([[f'"{c}"' if "," in c else c for c in row] for row in rows])


def _band_section(html: str) -> str:
    start = html.index('id="tabulasi"')
    end = html.index("</section>", start)
    return html[start:end]


def _stored_files(
    tabulasi_summary: str | None = None,
    tabulasi_item: str | None = None,
) -> dict[str, str]:
    files = {
        "item_table_15.1.csv": _csv([_ITEM_HEADER_ROW0, _ITEM_HEADER_ROW1, _make_item_row(1, "0.00", "I1")]),
        "option_table_15.3.csv": _csv([["ENTRY", "DATA"], ["", ""]]),
        "person_table.csv": _csv([_PERSON_HEADER_ROW0, _PERSON_HEADER_ROW1]),
        "summary_table.csv": _csv([["SECTION", "STATISTIC", "VALUE"], ["PERSON", "COUNT", "30"]]),
        "wright_map_measure.csv": _csv([_WRIGHT_HEADER_ROW0, _WRIGHT_HEADER_ROW1]),
        "wright_map_frequency.csv": _csv([_WRIGHT_HEADER_ROW0, _WRIGHT_HEADER_ROW1]),
    }
    if tabulasi_summary is not None:
        files["tabulasi_summary.csv"] = tabulasi_summary
    if tabulasi_item is not None:
        files["tabulasi_item.csv"] = tabulasi_item
    return files


def test_engine_run_stores_tabulasi_files(client: TestClient) -> None:
    raw = (FIXTURES_DIR / "sample_300x40.csv").read_bytes().decode()
    lines = raw.splitlines()
    header_cells = lines[0].split(",")

    new_header = [header_cells[0]]
    for i in range(1, 21):
        new_header.append(f"{i:02d}tbskd26a{i:02d}")
    for i in range(21, 41):
        new_header.append(f"{i:02d}tbspl25a{i:02d}")
    new_text = "\n".join([",".join(new_header)] + lines[1:]) + "\n"

    _create_authenticated_user(client, email="tabulasi_engine@example.test")
    resp = client.post(
        "/datasets",
        files={"data": ("tbs_sintetis.csv", new_text.encode(), "text/csv")},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    location = resp.headers["location"]
    dataset_id = int(location.split("/")[-1].split("?")[0])

    resp = client.post(
        f"/datasets/{dataset_id}/commit",
        data={
            "t0": "1",
            "m0": "correct",
            "t1": "0",
            "m1": "incorrect",
            "t2": "NA",
            "m2": "missing",
            "t3": "",
            "m3": "missing",
        },
        follow_redirects=False,
    )
    assert resp.status_code == 303

    resp = client.post(
        f"/datasets/{dataset_id}/analyze",
        data={"misfit": "1.50", "mode": "compat", "digits": "2", "person_order": "misfit"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    location = resp.headers["location"]
    analysis_id = int(location.split("/")[-1].split("?")[0])

    with SessionLocal() as db:
        row_summary = db.scalar(
            select(AnalysisFile).where(
                AnalysisFile.analysis_id == analysis_id,
                AnalysisFile.filename == "tabulasi_summary.csv",
            )
        )
        row_item = db.scalar(
            select(AnalysisFile).where(
                AnalysisFile.analysis_id == analysis_id,
                AnalysisFile.filename == "tabulasi_item.csv",
            )
        )
    assert row_summary is not None
    assert row_item is not None

    resp = client.get(f"/analyses/{analysis_id}")
    section = _band_section(resp.text)
    assert "Ringkasan Tabulasi" in section
    assert "SUBTES" in section
    assert "JUMLAH" in section
    assert "?table=tabulasi" in section
    assert "?table=tabulasi_butir" in section
    assert _OLD_RUN_NOTE not in section


def test_band_renders_summary_rows_and_download_links(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_render@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(
            tabulasi_summary=_csv_table([_TABULASI_HEADER] + _TABULASI_ROWS),
            tabulasi_item=_csv_table([_TABULASI_ITEM_HEADER] + _TABULASI_ITEM_ROWS),
        ),
    )
    resp = client.get(f"/analyses/{analysis_id}")
    section = _band_section(resp.text)

    for col in _TABULASI_HEADER:
        assert col in section
    assert "Aritmatika dan Aljabar" in section
    assert "Logis" in section
    assert "1.234" in section
    assert "1.237" in section
    assert "51, 54, 55" in section
    assert f'href="/analyses/{analysis_id}/export?table=tabulasi"' in section
    assert f'href="/analyses/{analysis_id}/export?table=tabulasi_butir"' in section
    assert _OLD_RUN_NOTE not in section
    assert _HEADER_ONLY_NOTE not in section


def test_band_notes_old_run_without_tabulasi_file(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_old@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(tabulasi_summary=None, tabulasi_item=None),
    )
    resp = client.get(f"/analyses/{analysis_id}")
    section = _band_section(resp.text)

    assert _OLD_RUN_NOTE in section
    assert "?table=tabulasi" not in section
    assert "?table=tabulasi_butir" not in section


def test_band_notes_header_only_file(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_header_only@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(
            tabulasi_summary=_csv_table([_TABULASI_HEADER]),
            tabulasi_item=_csv_table([_TABULASI_ITEM_HEADER] + _TABULASI_ITEM_ROWS),
        ),
    )
    resp = client.get(f"/analyses/{analysis_id}")
    section = _band_section(resp.text)

    assert _HEADER_ONLY_NOTE in section
    assert f'href="/analyses/{analysis_id}/export?table=tabulasi"' in section
    assert f'href="/analyses/{analysis_id}/export?table=tabulasi_butir"' in section


def test_band_sits_after_subsubtes_band(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_order@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(
            tabulasi_summary=_csv_table([_TABULASI_HEADER] + _TABULASI_ROWS),
            tabulasi_item=_csv_table([_TABULASI_ITEM_HEADER] + _TABULASI_ITEM_ROWS),
        ),
    )
    resp = client.get(f"/analyses/{analysis_id}")
    text = resp.text

    idx_subsubtes = text.index('id="subsubtes"')
    idx_tabulasi = text.index('id="tabulasi"')
    idx_tabel_butir = text.index("Tabel Butir (15.1)")
    assert idx_subsubtes < idx_tabulasi < idx_tabel_butir


def test_export_tabulasi_returns_workbook_with_sheet(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_export@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(
            tabulasi_summary=_csv_table([_TABULASI_HEADER] + _TABULASI_ROWS),
            tabulasi_item=_csv_table([_TABULASI_ITEM_HEADER] + _TABULASI_ITEM_ROWS),
        ),
    )
    resp = client.get(f"/analyses/{analysis_id}/export?table=tabulasi")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_MEDIA_TYPE

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    assert "tabulasi" in wb.sheetnames
    ws = wb["tabulasi"]

    first_row = [cell.value for cell in ws[1]]
    assert first_row == _TABULASI_HEADER

    second_row = [cell.value for cell in ws[2]]
    assert second_row[0] == "Kuantitatif"

    jumlah_idx = _TABULASI_HEADER.index("JUMLAH")
    jumlah_val = second_row[jumlah_idx]
    assert jumlah_val == 1237
    assert isinstance(jumlah_val, (int, float))
    assert not isinstance(jumlah_val, str)


def test_export_tabulasi_butir_and_semua_include_the_table(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_butir_semua@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(
            tabulasi_summary=_csv_table([_TABULASI_HEADER] + _TABULASI_ROWS),
            tabulasi_item=_csv_table([_TABULASI_ITEM_HEADER] + _TABULASI_ITEM_ROWS),
        ),
    )
    resp_butir = client.get(f"/analyses/{analysis_id}/export?table=tabulasi_butir")
    assert resp_butir.status_code == 200
    wb_butir = openpyxl.load_workbook(BytesIO(resp_butir.content))
    assert "tabulasi butir" in wb_butir.sheetnames
    ws_butir = wb_butir["tabulasi butir"]
    first_row = [cell.value for cell in ws_butir[1]]
    assert first_row == _TABULASI_ITEM_HEADER

    resp_semua = client.get(f"/analyses/{analysis_id}/export?table=semua")
    assert resp_semua.status_code == 200
    wb_semua = openpyxl.load_workbook(BytesIO(resp_semua.content))
    assert "tabulasi" in wb_semua.sheetnames


def test_export_absent_tabulasi_is_404(client: TestClient) -> None:
    user_id = _create_authenticated_user(client, email="tabulasi_absent@example.test")
    analysis_id = _seed_done(
        user_id,
        "tbs_ringkas.csv",
        _stored_files(tabulasi_summary=None, tabulasi_item=None),
    )
    resp = client.get(f"/analyses/{analysis_id}/export?table=tabulasi")
    assert resp.status_code == 404

    resp_butir = client.get(f"/analyses/{analysis_id}/export?table=tabulasi_butir")
    assert resp_butir.status_code == 404
